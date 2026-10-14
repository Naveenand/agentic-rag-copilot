"""Loaders, chunking, the knowledge base's sync logic, the audit log and settings."""

import socket

import pytest
from bs4 import BeautifulSoup
from conftest import SAMPLE_MD, SAMPLE_SOURCE

from app.core.config import Settings
from app.rag.retrieval import rrf_fuse
from app.services.audit import REDACTED, AuditLog
from app.services.ingestion import (
    IngestionError,
    check_public_url,
    chunk_markdown,
    clean_markdown,
    html_to_markdown,
    load_file,
    safe_filename,
    title_from_markdown,
)

# ---------------------------------------------------------------------------
# Cleaning and chunking
# ---------------------------------------------------------------------------


def test_clean_markdown_removes_boilerplate_and_front_matter():
    md = clean_markdown(SAMPLE_MD)
    assert "llms.txt" not in md and "title:" not in md


def test_chunks_are_section_aware_and_deterministic(settings):
    md, title = clean_markdown(SAMPLE_MD), title_from_markdown(SAMPLE_MD)
    chunks = chunk_markdown(md, SAMPLE_SOURCE, title, settings)
    sections = [c.metadata["section"] for c in chunks]
    assert any(s.endswith("Grade documents") for s in sections)
    assert not any("this comment" in s for s in sections)  # '#' inside code is not a heading
    assert all(c.page_content.startswith("Document: ") for c in chunks)
    assert [c.id for c in chunk_markdown(md, SAMPLE_SOURCE, title, settings)] == [c.id for c in chunks]


def test_html_fallback_skips_nav_and_duplicates():
    html = "<nav>MENU</nav><article><h2>Grade</h2><p>Para</p><ul><li>item <p>inner</p></li></ul><pre>code()</pre></article>"
    out = html_to_markdown(BeautifulSoup(html, "html.parser").find("article"))
    assert "MENU" not in out and out.count("inner") == 1 and "## Grade" in out


def test_rrf_rewards_documents_found_by_several_lists(kb):
    a, b, c = kb.store.all_documents()[:3]
    assert rrf_fuse([[a, b], [c, b], [b]])[0].id == b.id


# ---------------------------------------------------------------------------
# File loaders
# ---------------------------------------------------------------------------

LONG = "Employees must use the VPN before accessing internal systems. " * 4


def test_markdown_and_text_files(tmp_path):
    md = tmp_path / "vpn_guide.md"
    md.write_text(f"# VPN guide\n\n## Connect\n\n{LONG}")
    text, title = load_file(md)
    assert title == "VPN guide" and "## Connect" in text
    txt = tmp_path / "notes-2026.txt"
    txt.write_text(LONG)
    assert load_file(txt)[1] == "notes 2026"  # file name becomes the title


def test_docx_headings_become_sections(tmp_path, settings):
    import docx

    doc = docx.Document()
    doc.add_heading("IT Handbook", level=0)
    doc.add_heading("Password policy", level=1)
    doc.add_paragraph("Passwords must be at least 14 characters. " * 4)
    path = tmp_path / "handbook.docx"
    doc.save(path)
    md, title = load_file(path)
    assert title == "IT Handbook" and "## Password policy" in md
    chunks = chunk_markdown(md, "handbook.docx", title, settings)
    assert chunks[0].metadata["section"].endswith("Password policy")


def test_pdf_pages_become_sections(tmp_path, settings):
    from reportlab.pdfgen import canvas

    path = tmp_path / "policy.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(72, 720, "Page one text about multi-factor authentication for every employee account.")
    c.showPage()
    c.drawString(72, 720, "Page two text about laptop encryption and device management policies.")
    c.save()
    md, title = load_file(path)
    assert "## Page 1" in md and "## Page 2" in md and title == "policy"
    sections = {c.metadata["section"] for c in chunk_markdown(md, "policy.pdf", title, settings)}
    assert sections == {"Page 1", "Page 2"}


def test_unsupported_empty_and_corrupt_files_are_rejected(tmp_path):
    exe = tmp_path / "virus.exe"
    exe.write_bytes(b"MZ")
    with pytest.raises(IngestionError, match="Unsupported"):
        load_file(exe)
    empty = tmp_path / "empty.md"
    empty.write_text("   ")
    with pytest.raises(IngestionError, match="no extractable text"):
        load_file(empty)
    bad_pdf = tmp_path / "broken.pdf"
    bad_pdf.write_bytes(b"not a pdf")
    with pytest.raises(IngestionError, match="Could not read"):
        load_file(bad_pdf)


def test_safe_filename():
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("My Policy (v2).PDF") == "My_Policy_v2.pdf"
    assert safe_filename("") == "upload"


def test_private_urls_are_blocked(monkeypatch):
    with pytest.raises(IngestionError, match="private address"):
        check_public_url("http://127.0.0.1:8000/admin")
    with pytest.raises(IngestionError, match="Only http"):
        check_public_url("file:///etc/passwd")

    def fake_resolve(host, port):
        ip = "10.0.0.5" if host == "intranet.local" else "93.184.216.34"
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (ip, port))]

    monkeypatch.setattr(socket, "getaddrinfo", fake_resolve)
    with pytest.raises(IngestionError, match="private address"):
        check_public_url("https://intranet.local/wiki")
    check_public_url("https://docs.example.com/page")  # public → allowed


# ---------------------------------------------------------------------------
# Knowledge base sync
# ---------------------------------------------------------------------------


def test_reingesting_the_same_document_uploads_nothing(kb, vector_index):
    before = kb.store.count()
    r = kb.ingest(clean_markdown(SAMPLE_MD), source=SAMPLE_SOURCE, title="Build a custom RAG agent with LangGraph")
    assert r.uploaded == 0 and r.unchanged == r.chunks == before and r.removed == 0


def test_changed_document_removes_stale_chunks(kb, vector_index):
    old_ids = kb.store.ids_for_source(SAMPLE_SOURCE)
    new_md = (
        "# Build a custom RAG agent\n\n## Grade documents\n\n"
        + "Documents are now graded with a score from 1 to 5. " * 5
    )
    r = kb.ingest(new_md, source=SAMPLE_SOURCE, title="Build a custom RAG agent")
    assert r.removed == len(old_ids) and r.uploaded == r.chunks
    assert set(vector_index.deleted) == old_ids  # stale vectors gone from "Pinecone"
    assert kb.store.ids_for_source(SAMPLE_SOURCE) == set(vector_index.docs)
    assert all("score from 1 to 5" in d.page_content for d in kb.retriever.bm25_index.docs)  # BM25 refreshed


def test_new_document_is_searchable_immediately(kb):
    kb.ingest("# VPN guide\n\n" + LONG, source="vpn.md", title="VPN guide")
    hits = kb.retriever.retrieve(["How do employees use the VPN for internal systems?"])
    assert hits and hits[0].metadata["source"] == "vpn.md"
    assert {s["source"] for s in kb.store.sources()} == {SAMPLE_SOURCE, "vpn.md"}


def test_delete_source(kb, vector_index):
    n = kb.store.count()
    assert kb.delete_source(SAMPLE_SOURCE) == n
    assert kb.store.count() == 0 and not vector_index.docs and len(kb.retriever.bm25_index) == 0
    assert kb.delete_source("missing") == 0


def test_tiny_document_is_rejected(kb):
    with pytest.raises(IngestionError, match="no chunks"):
        kb.ingest("# Title\n\nToo short.", source="tiny.md", title="Tiny")


# ---------------------------------------------------------------------------
# Audit log and settings
# ---------------------------------------------------------------------------


def test_audit_log_redacts_private_questions(tmp_path):
    log = AuditLog(tmp_path / "audit.db")
    log.record(
        question="What is RAG?",
        request_id="r1",
        thread_id="t",
        result={"route": "retrieve", "source_used": "private_kb", "latency_s": 2.0, "trace": ["a"]},
    )
    log.record(question="Customer account 88231 is locked", request_id="r2", thread_id="t", error="RuntimeError")
    rows = log.recent()
    assert rows[0]["question"] == REDACTED and rows[0]["error"] == "RuntimeError"
    assert rows[1]["question"] == "What is RAG?" and rows[1]["trace"] == ["a"]
    assert log.summary() == {"questions": 2, "avg_latency_s": 2.0, "by_source": {None: 1, "private_kb": 1}, "errors": 1}


def test_production_requires_a_strong_admin_key():
    with pytest.raises(RuntimeError, match="ADMIN_API_KEY"):
        Settings(_env_file=None, app_env="production").check_production_safety()
    Settings(_env_file=None, app_env="production", admin_api_key="x" * 32).check_production_safety()


def test_missing_keys_ignores_tavily_when_web_is_off():
    s = Settings(_env_file=None, openai_api_key="k", pinecone_api_key="k", tavily_api_key="", enable_web=False)
    assert s.missing_keys() == []
    s2 = Settings(_env_file=None, openai_api_key="", pinecone_api_key="k", tavily_api_key="k")
    assert s2.missing_keys() == ["OPENAI_API_KEY"]


def test_keyword_index_picks_up_changes_from_another_process(kb, settings, vector_index):
    from app.services.knowledge_base import ChunkStore, KnowledgeBase

    assert kb.refresh_if_stale() is False  # nothing changed since the fixture's ingest
    cli_kb = KnowledgeBase(settings, ChunkStore(settings.chunk_db_path), vector_index, retriever=None)  # like cli.py
    cli_kb.ingest("# VPN guide\n\n" + LONG, source="vpn.md", title="VPN guide")
    assert all(d.metadata["source"] != "vpn.md" for d in kb.retriever.bm25_index.docs)  # server not told yet
    assert kb.refresh_if_stale() is True
    assert any(d.metadata["source"] == "vpn.md" for d in kb.retriever.bm25_index.docs)
